"""Native Codex web handoff, grounded notes and portable local reports."""

import html
import json
import uuid
import xml.etree.ElementTree as ET

from companion_codex.native import public_url

from .cognition import stable_id
from .errors import AssistantError
from .knowledge import annotations, ref, source_for
from .storage import atomic_write


def capture(engine, url, title, excerpt, tool_reference):
    """The foreground supplies actual native-web output; no claim of independent HTTP verification."""
    engine.mode("role")
    url = public_url(url)
    if not excerpt.strip() or not tool_reference.strip():
        raise AssistantError("web_material", "需要实际网页摘录及 Codex 原生读取引用，搜索标题不算材料")
    account = engine.store.read("manifest")["data"]["account_uid"]
    item_id = stable_id("web", json.dumps([account, url, excerpt], ensure_ascii=False))
    try:
        return {"duplicate": True, **engine.store.read("source_snapshots", item_id)}
    except AssistantError as exc:
        if exc.code != "not_found":
            raise
    return engine.store.write("source_snapshots", {
        "schema_version": 1, "id": item_id, "account_uid": account, "bvid": "", "url": url,
        "title": title, "owner_uid": "", "visual_note": "",
        "coverage": {"kind": "web", "bvid": "", "start_seconds": 0, "end_seconds": 0,
                     "duration_seconds": 0, "complete": False},
        "segments": [{"start": 0, "end": 0, "text": excerpt}],
        "capture": {"method": "codex_native_web", "tool_reference": tool_reference, "scope": "excerpt"},
        "source": {"type": "web", "recorded_at": engine.clock().isoformat(), "reference": url}, "locked": False,
    }, None, engine.identity(), item_id)


def save_note(engine, source_id, source_version, content, quotes, tags):
    engine.mode("role")
    source = engine.store.read("source_snapshots", source_id)
    if (source["version"] != source_version or source["data"]["coverage"]["kind"] != "web"
            or source["data"]["account_uid"] != engine.store.read("manifest")["data"]["account_uid"]):
        raise AssistantError("context_changed", "网页材料已变化或不属于当前账号")
    if not quotes or any(not q.strip() or not any(q in s["text"] for s in source["data"]["segments"]) for q in quotes):
        raise AssistantError("unverified_quote", "网页笔记需要逐字存在于已保存材料中的引用")
    item_id = stable_id("knowledge-web", source_version + content)
    return engine.remember("knowledge", {
        "schema_version": 1, "id": item_id, "content": content, "source": source["data"]["source"],
        "source_ref": ref(source), "coverage": source["data"]["coverage"],
        "source_quotes": quotes,
        "tags": tags, "locked": False, "disagrees_with": [],
    })


def attach(engine, goal_id, knowledge_id, expected_version):
    from .learning import status

    engine.mode("role")
    with engine.store.lock:
        goal = status(engine, goal_id)["goals"][0]
        if not goal["request_current"] or goal["data"]["status"] == "cancelled":
            raise AssistantError("goal_finished", "目标要求已变化或已取消")
        note = engine.store.read("knowledge", knowledge_id)
        if not source_for(engine.store, note, goal["data"]["account_uid"]):
            raise AssistantError("unverified_source", "知识没有当前账号可复核的来源")
        data = goal["data"]
        refs = [r for r in data["knowledge_refs"] if r["id"] != knowledge_id]
        if len(refs) >= 20:
            raise AssistantError("goal_material_limit", "目标材料已满，请先整理本地引用")
        data.update(knowledge_refs=[*refs, ref(note)], phase="quiz", reason="foreground_material_attached",
                    updated_at=engine.clock().isoformat())
        if data["status"] == "completed":
            data["status"] = "paused"
        return engine.store.write("learning_goals", data, expected_version, engine.identity(), goal_id)


def export(engine, title, knowledge_ids, format="markdown", style="report", goal_id=None):
    engine.mode("role")
    store = engine.store
    account = store.read("manifest")["data"]["account_uid"]
    notes = annotations(store, [{"kind": "knowledge", **store.read("knowledge", key)}
                               for key in dict.fromkeys(knowledge_ids)])
    from .learning_tree import status

    tree = status(engine, goal_id) if goal_id else None
    if not notes and not tree:
        raise AssistantError("report_empty", "至少选择一条知识笔记或一个学习目标")
    references, sections, stale = [], [], []
    for index, note in enumerate(notes, 1):
        doc = note["data"]
        source = source_for(store, note, account)
        references.append({"kind": "knowledge", **ref(note)})
        label = f"{'知识卡' if style == 'cards' else '材料'} {index} · {' / '.join(doc['tags'][:3]) or doc['id']}"
        body = doc["content"]
        source_text = "来源缺失或版本已变化，待重新核验。"
        if source:
            references.append({"kind": "source_snapshots", **ref(source)})
            coverage = source["data"]["coverage"]
            scope = "网页摘录（由角色聊天原生 web 工具导入）" if coverage["kind"] == "web" else (
                f"{coverage['kind']}，{coverage['start_seconds']:g}–{coverage['end_seconds']:g} 秒，"
                + ("材料覆盖完整" if coverage["complete"] else "部分覆盖"))
            source_text = f"{source['data']['title']}\n{source['data']['url']}\n范围：{scope}"
        else:
            stale.append(doc["id"])
        review = note.get("review")
        review_text = review["data"]["result"]["summary"] if review else "尚无当前版本的质量复查。"
        sections.append((label, f"{body}\n\n来源：{source_text}\n复查：{review_text}"))
    if tree:
        references.append({"kind": "learning_goals", **ref(tree["goal"])})
        sections.append(("学习目标", tree["goal"]["data"]["objective"]))
        for node in tree["nodes"]:
            doc = node["data"]
            references.append({"kind": "learning_nodes", **ref(node)})
            score = f"{doc['mastery']:.0%}" if node["evidence_current"] and doc["mastery"] is not None else "待核验"
            mistakes = "\n".join(m["gap"] for m in doc["mistakes"] if not m["resolved"]) or "无未解决错题"
            sections.append((doc["title"], f"模型自检掌握度：{score}\n下次复习：{doc['next_review_at'] or '未安排'}\n{mistakes}"))
    notice = "由本地材料整理；模型自检不是独立实测。缺失、部分覆盖与过期依据已按实际状态标注。"
    directory = (store.root / "exports").resolve()
    if not directory.is_relative_to(store.root):
        raise AssistantError("invalid_path", "导出目录超出本地数据目录")
    directory.mkdir(parents=True, exist_ok=True)
    stem = "report-" + uuid.uuid4().hex[:16]
    suffix = {"markdown": ".md", "html": ".html", "docx": ".docx", "mindmap": ".mm"}[format]
    target = directory / (stem + suffix)
    if format == "markdown":
        text = f"# {title}\n\n{notice}\n\n" + "\n\n".join(f"## {name}\n\n{body}" for name, body in sections)
        atomic_write(target, text.encode("utf-8"))
    elif format == "html":
        # No raw source markup, remote scripts, tracking pixels or untrusted CSS.
        text = '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>' + html.escape(title) + '</title>'
        text += '<style>body{max-width:900px;margin:40px auto;padding:0 24px;font:17px/1.7 system-ui;background:#faf9f6;color:#242424}details{border:1px solid #ddd;border-radius:12px;padding:18px;margin:18px 0}pre{white-space:pre-wrap;font:inherit;overflow-wrap:anywhere}summary{font-weight:600;cursor:pointer}@media print{details{break-inside:avoid}}</style>'
        text += f"<h1>{html.escape(title)}</h1><p>{notice}</p>"
        text += "".join(f"<details open><summary>{html.escape(name)}</summary><pre>{html.escape(body)}</pre></details>"
                        for name, body in sections)
        atomic_write(target, (text + "</html>").encode("utf-8"))
    elif format == "docx":
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = Document()
        document.add_heading(title, 0)
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        document.styles["Normal"].element.get_or_add_rPr().append(fonts)
        document.add_paragraph(notice)
        for name, body in sections:
            document.add_heading(name, 1)
            for paragraph in body.split("\n"):
                document.add_paragraph(paragraph)
        document.save(target)
    else:
        root = ET.Element("map", version="1.0.1")
        center = ET.SubElement(root, "node", TEXT=title)
        if tree:
            elements = {n["data"]["id"]: ET.Element("node", TEXT=n["data"]["title"]) for n in tree["nodes"]}
            for node in tree["nodes"]:
                doc = node["data"]
                parent = elements.get(doc["parent_id"]) if not tree["graph_error"] else None
                (parent if parent is not None else center).append(elements[doc["id"]])
                for pre in doc["prerequisites"]:
                    ET.SubElement(elements[doc["id"]], "node", TEXT="先修：" + pre)
        for name, body in sections[:len(notes)]:
            item = ET.SubElement(center, "node", TEXT=name)
            ET.SubElement(item, "node", TEXT=body)
        atomic_write(target, ET.tostring(root, encoding="utf-8", xml_declaration=True))
    manifest = {"title": title, "format": format, "style": style, "created_at": engine.clock().isoformat(),
                "references": references, "stale_knowledge_ids": stale, "path": str(target)}
    atomic_write(directory / (stem + ".json"), json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8"))
    return manifest
